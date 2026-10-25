"""Generate a fictional Word note and a fictional plain-text note, the two
other ways documents reach the app (Word export, copy-paste from the
patient record). Every person and identifier below is invented.

The Word note keeps identifiers in a table, as admission forms do. The text
note is saved as Windows-1252 with CRLF line ends, as older exports are.
Their planted identifiers are added to ``ground_truth.json`` next to the
PDFs' (which make_synthetic_pdfs.py writes).

    uv run python scripts/make_synthetic_texts.py
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document

OUT_DIR = Path(__file__).resolve().parent.parent / "data" / "anonymized"
TRUTH = OUT_DIR / "ground_truth.json"


def word_note() -> list[str]:
    doc = Document()
    doc.add_heading("Fiche de liaison - Consultation addictologie", level=1)
    table = doc.add_table(rows=0, cols=2)
    for label, value in [
        ("Nom", "BOUVIER"),
        ("Prénom", "Kévin"),
        ("Né le", "03/02/1991"),
        ("IPP", "0078451236"),
        ("Adresse", "7 impasse des Mésanges, 33600 PESSAC"),
        ("Téléphone", "07 81 22 45 90"),
    ]:
        row = table.add_row().cells
        row[0].text, row[1].text = label, value
    doc.add_paragraph(
        "Consultation du 18/09/2024. Patient de 33 ans adressé par le Dr GIRAUDEAU pour mésusage d'alcool. "
        "Vit avec sa compagne Mme Laetitia PRUVOST. Consommation déclarée : 8 à 10 "
        "verres/jour, majoration depuis la perte d'emploi."
    )
    doc.add_paragraph(
        "Conduite à tenir : sevrage ambulatoire, diazépam selon protocole, revoir dans 8 jours. "
        "Courrier au médecin traitant, Dr GIRAUDEAU."
    )
    doc.save(OUT_DIR / "synthetique-fiche-word.docx")
    return ["BOUVIER", "Kévin", "03/02/1991", "0078451236", "7 impasse des Mésanges",
            "07 81 22 45 90", "GIRAUDEAU", "PRUVOST", "Laetitia"]


def text_note() -> list[str]:
    text = "\n".join([
        "12/03/2024 - IDE - CHAPUIS Hélène",
        "Mr RIBEIRO Fabrice né le 21/11/1969, IPP 0061237845.",
        "Entrée pour sevrage alcool. Arrivé à 9h, TA 14/9, tremblements +, Cushman 7.",
        "Appel de sa soeur Mme Sandrine RIBEIRO au 06 45 12 78 33, demande des nouvelles.",
        "",
        "13/03/2024 - Dr LEBRETON",
        "Cushman 3. Poursuite décroissance oxazépam. Sortie envisagée vendredi.",
    ])
    (OUT_DIR / "synthetique-note-collee.txt").write_bytes(text.replace("\n", "\r\n").encode("cp1252"))
    return ["CHAPUIS", "Hélène", "RIBEIRO", "Fabrice", "21/11/1969", "0061237845",
            "Sandrine", "06 45 12 78 33", "LEBRETON"]


def main() -> None:
    truth = json.loads(TRUTH.read_text())
    truth["synthetique-fiche-word.docx"] = word_note()
    truth["synthetique-note-collee.txt"] = text_note()
    TRUTH.write_text(json.dumps(truth, ensure_ascii=False, indent=2) + "\n")
    print(f"wrote 2 documents to {OUT_DIR}")


if __name__ == "__main__":
    main()
